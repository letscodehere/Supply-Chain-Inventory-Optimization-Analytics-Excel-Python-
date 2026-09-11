import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DATA_DIR = "/home/claude/supplychain_project/data"
OUT_DIR = "/home/claude/supplychain_project/outputs"
XLSX_PATH = "/home/claude/supplychain_project/excel/Supply_Chain_Inventory_Optimization.xlsx"

NAVY, TEAL, GOLD, RED = "1F3864", "2E8B8B", "D4A017", "B23A48"
LIGHT, YELLOW, WHITE = "EDF1F7", "FFF2CC", "FFFFFF"
HEADER_FONT = Font(name="Arial", bold=True, color=WHITE, size=11)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(name="Arial", bold=True, color=NAVY, size=18)
SUBTITLE_FONT = Font(name="Arial", italic=True, color="595959", size=10)
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", italic=True, size=9, color="808080")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YELLOW_FILL = PatternFill("solid", fgColor=YELLOW)
INPUT_FILL = PatternFill("solid", fgColor="D9E8F5")
INPUT_FONT = Font(name="Arial", bold=True, color="1F3864", size=11)

abc = pd.read_csv(f"{OUT_DIR}/abc_analysis.csv")
abc_summary = pd.read_csv(f"{OUT_DIR}/abc_summary.csv")
reorder = pd.read_csv(f"{OUT_DIR}/reorder_parameters.csv")

wb = load_workbook(XLSX_PATH)


def style_header_row(ws, row, ncols, start_col=1):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def write_small_table(ws, df, start_row, start_col=1, table_name=None):
    for j, col in enumerate(df.columns):
        ws.cell(row=start_row, column=start_col + j, value=col)
    style_header_row(ws, start_row, len(df.columns), start_col)
    r0 = start_row + 1
    for i, row in enumerate(df.itertuples(index=False)):
        for j, val in enumerate(row):
            cell = ws.cell(row=r0 + i, column=start_col + j, value=val)
            cell.font = BODY_FONT
            cell.border = BORDER
    end_row = r0 + len(df) - 1
    if len(df) > 0:
        end_col_letter = get_column_letter(start_col + len(df.columns) - 1)
        start_col_letter = get_column_letter(start_col)
        ref = f"{start_col_letter}{start_row}:{end_col_letter}{end_row}"
        tbl = Table(displayName=table_name, ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)
    return end_row


def autosize(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# ====================================================================
# SHEET — ABC Analysis
# ====================================================================
ws = wb.create_sheet("ABC Analysis")
ws.sheet_view.showGridLines = False
ws["B2"] = "ABC Analysis — SKU Consumption Value"
ws["B2"].font = TITLE_FONT
ws["B3"] = "Precomputed in Python from Data (ranking + cumulative sums across 20 SKUs)."
ws["B3"].font = SUBTITLE_FONT

r = 5
ws.cell(row=r, column=2, value="ABC Summary").font = Font(name="Arial", bold=True, size=12, color=NAVY)
end1 = write_small_table(ws, abc_summary, r + 1, start_col=2, table_name="ABCSummary")
for rr in range(r + 2, end1 + 1):
    for cc in (2, 3, 4):
        ws.cell(row=rr, column=cc).fill = YELLOW_FILL
    ws.cell(row=rr, column=5).number_format = '0.0"%"'   # sku_share_pct, already 0-100 scale
    ws.cell(row=rr, column=5).fill = YELLOW_FILL
    ws.cell(row=rr, column=6).number_format = '0.0"%"'   # value_share_pct, already 0-100 scale
    ws.cell(row=rr, column=6).fill = YELLOW_FILL

r2 = end1 + 3
ws.cell(row=r2, column=2, value=(
    "Note: unlike many real-world retail portfolios (which often show 20% of SKUs driving "
    "70-80% of value), this dataset's consumption value is close to evenly spread across SKUs "
    "(A-class is 65% of SKUs by count here, not the classic ~20%). The standard 70/90 cumulative-% "
    "cutoffs are still applied for methodological completeness — see the business report for detail."
)).font = NOTE_FONT
ws.cell(row=r2, column=2).alignment = Alignment(wrap_text=True)
ws.row_dimensions[r2].height = 45

r3 = r2 + 3
ws.cell(row=r3, column=2, value="Full SKU Ranking").font = Font(name="Arial", bold=True, size=12, color=NAVY)
abc_display = abc[["rank", "sku", "category", "total_consumption_value", "value_share_pct",
                    "cumulative_value_pct", "abc_class"]]
end2 = write_small_table(ws, abc_display, r3 + 1, start_col=2, table_name="ABCRanking")
for rr in range(r3 + 2, end2 + 1):
    ws.cell(row=rr, column=5).number_format = '$#,##0'
    ws.cell(row=rr, column=6).number_format = '0.0"%"'   # values already in 0-100 scale
    ws.cell(row=rr, column=7).number_format = '0.0"%"'   # values already in 0-100 scale
    for cc in range(2, 9):
        ws.cell(row=rr, column=cc).fill = YELLOW_FILL

autosize(ws, {"A": 2, "B": 10, "C": 10, "D": 14, "E": 18, "F": 14, "G": 16, "H": 12})

chart = BarChart()
chart.type = "col"
chart.title = "SKU Consumption Value Ranking"
data_ref = Reference(ws, min_col=6, min_row=r3 + 1, max_row=end2)
cats_ref = Reference(ws, min_col=3, min_row=r3 + 2, max_row=end2)
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
chart.height, chart.width = 9, 18
ws.add_chart(chart, "J5")

print("ABC Analysis sheet written.")

# ====================================================================
# SHEET — Safety Stock & Reorder Point Calculator (LIVE FORMULAS)
# ====================================================================
ws = wb.create_sheet("Safety Stock and ROP Calculator")
ws.sheet_view.showGridLines = False
ws["B2"] = "Safety Stock & Reorder Point Calculator"
ws["B2"].font = TITLE_FONT
ws["B3"] = "Live formulas — change any blue assumption cell and every row recalculates."
ws["B3"].font = SUBTITLE_FONT

ws.cell(row=5, column=2, value="Assumptions (editable)").font = Font(name="Arial", bold=True, size=12, color=NAVY)
assumptions = [
    ("Target Service Level", 0.95, '0%'),
    ("Z-score (from service level)", "=NORMSINV(C6)", '0.00'),
    ("Ordering Cost per Order ($)", 50, '$#,##0'),
    ("Holding Cost (% of unit value / year)", 0.20, '0%'),
]
for i, (label, val, fmt) in enumerate(assumptions):
    rr = 6 + i
    ws.cell(row=rr, column=2, value=label).font = BODY_FONT
    cell = ws.cell(row=rr, column=3, value=val)
    cell.font = INPUT_FONT
    cell.fill = INPUT_FILL
    cell.number_format = fmt
    cell.border = BORDER
    ws.cell(row=rr, column=2).border = BORDER

ws.cell(row=11, column=2, value=(
    "Blue cells are inputs. Z-score updates automatically from the service-level target via "
    "NORM.S.INV. Lead times (below) are a stated per-category assumption — see Read Me and the "
    "business report; substitute real supplier lead times here for a production version."
)).font = NOTE_FONT
ws.cell(row=11, column=2).alignment = Alignment(wrap_text=True)
ws.row_dimensions[11].height = 30

start = 14
ws.cell(row=start - 1, column=2, value="Per Warehouse-SKU Calculation").font = Font(
    name="Arial", bold=True, size=12, color=NAVY)

headers = ["Warehouse", "SKU", "Category", "Mean Daily Demand", "Std Dev Daily Demand",
           "Lead Time (days)", "Safety Stock", "Reorder Point", "Avg. Unit Price", "EOQ"]
for j, h in enumerate(headers):
    ws.cell(row=start, column=2 + j, value=h)
style_header_row(ws, start, len(headers), 2)

for i, row in enumerate(reorder.itertuples(index=False)):
    rr = start + 1 + i
    ws.cell(row=rr, column=2, value=row.warehouse_id).font = BODY_FONT
    ws.cell(row=rr, column=3, value=row.sku).font = BODY_FONT
    ws.cell(row=rr, column=4, value=row.category).font = BODY_FONT
    d_cell = ws.cell(row=rr, column=5, value=round(row.mean_daily_demand, 1))
    d_cell.number_format = '0.0'
    s_cell = ws.cell(row=rr, column=6, value=round(row.std_daily_demand, 1))
    s_cell.number_format = '0.0'
    lt_cell = ws.cell(row=rr, column=7, value=int(row.lead_time_days))
    lt_cell.number_format = '0'
    ss_cell = ws.cell(row=rr, column=8, value=f"=ROUND($C$7*F{rr}*SQRT(G{rr}),0)")
    ss_cell.number_format = '#,##0'
    rop_cell = ws.cell(row=rr, column=9, value=f"=ROUND(E{rr}*G{rr}+H{rr},0)")
    rop_cell.number_format = '#,##0'
    price_cell = ws.cell(row=rr, column=10, value=round(row.mean_price, 2))
    price_cell.number_format = '$#,##0.00'
    eoq_cell = ws.cell(row=rr, column=11, value=f"=ROUND(SQRT(2*E{rr}*365*$C$8/(J{rr}*$C$9)),0)")
    eoq_cell.number_format = '#,##0'
    for cc in range(2, 12):
        ws.cell(row=rr, column=cc).border = BORDER

end_calc = start + len(reorder)
autosize(ws, {"A": 2, "B": 11, "C": 8, "D": 12, "E": 15, "F": 16, "G": 12, "H": 12,
               "I": 13, "J": 12, "K": 10})

wb.save(XLSX_PATH)
print("Safety Stock & ROP Calculator sheet written.")
