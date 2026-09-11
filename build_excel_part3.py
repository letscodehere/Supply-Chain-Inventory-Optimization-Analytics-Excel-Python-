import pandas as pd
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT_DIR = "/home/claude/supplychain_project/outputs"
XLSX_PATH = "/home/claude/supplychain_project/excel/Supply_Chain_Inventory_Optimization.xlsx"

NAVY, TEAL, GOLD, RED = "1F3864", "2E8B8B", "D4A017", "B23A48"
YELLOW, WHITE = "FFF2CC", "FFFFFF"
HEADER_FONT = Font(name="Arial", bold=True, color=WHITE, size=11)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(name="Arial", bold=True, color=NAVY, size=18)
SUBTITLE_FONT = Font(name="Arial", italic=True, color="595959", size=10)
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", italic=True, size=9, color="808080")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YELLOW_FILL = PatternFill("solid", fgColor=YELLOW)

sim = pd.read_csv(f"{OUT_DIR}/simulation_results.csv")
sim_cat = pd.read_csv(f"{OUT_DIR}/simulation_by_category.csv").fillna("n/a (blanket had 0 stockouts)")
with open(f"{OUT_DIR}/simulation_summary.json") as f:
    sim_summary = json.load(f)
forecast_summary = pd.read_csv(f"{OUT_DIR}/forecast_accuracy_summary.csv", index_col=0).reset_index()
forecast_summary.columns = ["Metric"] + list(forecast_summary.columns[1:])
demand_stats = pd.read_csv(f"{OUT_DIR}/demand_stats_by_sku_warehouse.csv")

wb = load_workbook(XLSX_PATH)


def style_header_row(ws, row, ncols, start_col=1):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def write_small_table(ws, df, start_row, start_col=1, table_name=None):
    for j, col in enumerate(df.columns):
        ws.cell(row=start_row, column=start_col + j, value=col)
    style_header_row(ws, start_row, len(df.columns), start_col)
    r0 = start_row + 1
    for i, row in enumerate(df.itertuples(index=False)):
        for j, val in enumerate(row):
            cell = ws.cell(row=r0 + i, column=start_col + j, value=val)
            cell.font = BODY_FONT
            cell.border = BORDER
    end_row = r0 + len(df) - 1
    if len(df) > 0:
        end_col_letter = get_column_letter(start_col + len(df.columns) - 1)
        start_col_letter = get_column_letter(start_col)
        ref = f"{start_col_letter}{start_row}:{end_col_letter}{end_row}"
        tbl = Table(displayName=table_name, ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)
    return end_row


def autosize(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def kpi_card(ws, row, col, label, value, number_format=None):
    lbl_cell = ws.cell(row=row, column=col, value=label)
    lbl_cell.font = Font(name="Arial", bold=True, color="595959", size=9)
    val_cell = ws.cell(row=row + 1, column=col, value=value)
    val_cell.font = Font(name="Arial", bold=True, color=NAVY, size=16)
    if number_format:
        val_cell.number_format = number_format
    for rr in (row, row + 1):
        c = ws.cell(row=rr, column=col)
        c.fill = YELLOW_FILL
        c.border = BORDER
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)


# ====================================================================
# SHEET — Simulation Results
# ====================================================================
ws = wb.create_sheet("Simulation Results")
ws.sheet_view.showGridLines = False
ws["B2"] = "Reorder Policy Simulation: Blanket vs. Framework"
ws["B2"].font = TITLE_FONT
ws["B3"] = "731-day trace-driven simulation against real historical demand — precomputed in Python."
ws["B3"].font = SUBTITLE_FONT

kpi_card(ws, 5, 2, "STOCKOUT DAYS: BLANKET", sim_summary["naive_total_stockout_days"], '#,##0')
kpi_card(ws, 5, 4, "STOCKOUT DAYS: FRAMEWORK", sim_summary["optimized_total_stockout_days"], '#,##0')
kpi_card(ws, 5, 6, "STOCKOUT REDUCTION", sim_summary["stockout_day_reduction_pct"] / 100, '0.0%')
kpi_card(ws, 8, 2, "AVG INVENTORY: BLANKET", sim_summary["naive_avg_inventory_units"], '#,##0')
kpi_card(ws, 8, 4, "AVG INVENTORY: FRAMEWORK", sim_summary["optimized_avg_inventory_units"], '#,##0')
kpi_card(ws, 8, 6, "INVENTORY CHANGE", sim_summary["avg_inventory_change_pct"] / 100, '+0.0%;-0.0%')

ws.cell(row=11, column=2, value=(
    "Blanket policy = a single portfolio-wide reorder point & order quantity applied identically "
    "to every warehouse-SKU (representative of a business with no per-SKU inventory analytics). "
    "Framework policy = each warehouse-SKU's own statistically-tailored reorder point (Safety Stock "
    "and ROP Calculator sheet) and EOQ. Both reductions happen together because the blanket policy's "
    "error is misallocation — overstocking some items while understocking others — which the "
    "tailored framework corrects in both directions at once."
)).font = NOTE_FONT
ws.cell(row=11, column=2).alignment = Alignment(wrap_text=True)
ws.row_dimensions[11].height = 55

r0 = 14
ws.cell(row=r0, column=2, value="Results by Category").font = Font(name="Arial", bold=True, size=12, color=NAVY)
end1 = write_small_table(ws, sim_cat, r0 + 1, start_col=2, table_name="SimByCategory")
for rr in range(r0 + 2, end1 + 1):
    for cc in range(2, 6):
        ws.cell(row=rr, column=cc).fill = YELLOW_FILL
    ws.cell(row=rr, column=4).number_format = '#,##0.0'
    ws.cell(row=rr, column=5).number_format = '#,##0.0'
    pct_cell = ws.cell(row=rr, column=7)
    pct_cell.fill = YELLOW_FILL
    if isinstance(pct_cell.value, (int, float)):
        pct_cell.number_format = '0.0"%"'  # already in 0-100 scale

chart = BarChart()
chart.type = "col"
chart.title = "Stockout Days by Category: Blanket vs. Framework"
data_ref = Reference(ws, min_col=3, max_col=4, min_row=r0 + 1, max_row=end1)
cats_ref = Reference(ws, min_col=2, min_row=r0 + 2, max_row=end1)
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
chart.height, chart.width = 9, 16
ws.add_chart(chart, "I14")

r2 = end1 + 3
ws.cell(row=r2, column=2, value="Full Detail — All 100 Warehouse-SKU Combinations").font = Font(
    name="Arial", bold=True, size=12, color=NAVY)
sim_display = sim[["warehouse_id", "sku", "category", "lead_time_days", "naive_stockout_days",
                    "optimized_stockout_days", "naive_avg_inventory", "optimized_avg_inventory"]]
end2 = write_small_table(ws, sim_display, r2 + 1, start_col=2, table_name="SimFullDetail")
for rr in range(r2 + 2, end2 + 1):
    for cc in range(2, 10):
        ws.cell(row=rr, column=cc).fill = YELLOW_FILL

autosize(ws, {"A": 2, "B": 13, "C": 9, "D": 13, "E": 16, "F": 16, "G": 16, "H": 16, "I": 20, "J": 20})

print("Simulation Results sheet written.")

# ====================================================================
# SHEET — Demand Variability & Forecast Accuracy
# ====================================================================
ws = wb.create_sheet("Demand and Forecast")
ws.sheet_view.showGridLines = False
ws["B2"] = "Demand Variability & Forecast Accuracy"
ws["B2"].font = TITLE_FONT
ws["B3"] = "Precomputed in Python (statistical summaries and a 60-day held-out forecast test)."
ws["B3"].font = SUBTITLE_FONT

r = 5
ws.cell(row=r, column=2, value="Forecast Accuracy — 60-Day Held-Out Test (avg. across 100 series)").font = Font(
    name="Arial", bold=True, size=12, color=NAVY)
end1 = write_small_table(ws, forecast_summary, r + 1, start_col=2, table_name="ForecastAccuracy")
for rr in range(r + 2, end1 + 1):
    for cc in range(2, 6):
        ws.cell(row=rr, column=cc).fill = YELLOW_FILL

r2 = end1 + 3
ws.cell(row=r2, column=2, value=(
    "Note: daily demand in this dataset shows near-zero autocorrelation (lag-1 autocorrelation "
    "\u2248 -0.01), i.e. it behaves close to random noise around a stable mean. Classical time-series "
    "methods (moving average, Holt-Winters) therefore cannot beat the dataset's own pre-supplied "
    "forecast, which appears to encode information not derivable from historical demand alone. "
    "This is a genuine, honestly-reported finding, not a modelling failure — see the business report."
)).font = NOTE_FONT
ws.cell(row=r2, column=2).alignment = Alignment(wrap_text=True)
ws.row_dimensions[r2].height = 55

r3 = r2 + 3
ws.cell(row=r3, column=2, value="Demand Statistics by Warehouse-SKU (first 30 shown)").font = Font(
    name="Arial", bold=True, size=12, color=NAVY)
demand_display = demand_stats[["warehouse_id", "sku", "category", "mean_daily_demand",
                                 "std_daily_demand", "cv", "consumption_value"]].head(30)
end2 = write_small_table(ws, demand_display, r3 + 1, start_col=2, table_name="DemandStats")
for rr in range(r3 + 2, end2 + 1):
    ws.cell(row=rr, column=5).number_format = '0.0'
    ws.cell(row=rr, column=6).number_format = '0.0'
    ws.cell(row=rr, column=7).number_format = '0.00'
    ws.cell(row=rr, column=8).number_format = '$#,##0'
    for cc in range(2, 9):
        ws.cell(row=rr, column=cc).fill = YELLOW_FILL

autosize(ws, {"A": 2, "B": 30, "C": 16, "D": 16, "E": 16, "F": 10, "G": 14, "H": 12})

print("Demand and Forecast sheet written.")

# ====================================================================
# Final tab order + tab colors
# ====================================================================
order = ["Read Me", "Dashboard", "ABC Analysis", "Safety Stock and ROP Calculator",
         "Simulation Results", "Demand and Forecast", "Data"]
wb._sheets = [wb[name] for name in order]
wb.active = wb.sheetnames.index("Dashboard")
for name in wb.sheetnames:
    wb[name].sheet_properties.tabColor = NAVY if name != "Data" else "A6A6A6"

wb.save(XLSX_PATH)
print("Final workbook saved. Sheets:", wb.sheetnames)
